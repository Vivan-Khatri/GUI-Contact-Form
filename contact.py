import re
import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk

from openpyxl import Workbook, load_workbook


FILE_PATH = Path(__file__).with_name("contact_data.xlsx")
SHEET_NAME = "Contacts"
HEADERS = ["Name", "Email", "Mobile No", "Address", "Branch"]

NAME_PATTERN = re.compile(r"^[A-Za-z][A-Za-z\s]{1,49}$")
EMAIL_PATTERN = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
MOBILE_PATTERN = re.compile(r"^\d{10}$")
ADDRESS_PATTERN = re.compile(r"^[A-Za-z0-9\s,./#-]{5,100}$")
BRANCH_PATTERN = re.compile(r"^[A-Za-z0-9\s-]{2,30}$")


def get_workbook_and_sheet():
    # Open the workbook if it exists; otherwise create it with a header row.
    if FILE_PATH.exists():
        workbook = load_workbook(FILE_PATH)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
        sheet.title = SHEET_NAME
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(HEADERS)

    return workbook, sheet


def save_workbook(workbook):
    try:
        workbook.save(FILE_PATH)
        return True, f"Saved to {FILE_PATH.name}"
    except PermissionError:
        return False, f"Could not save {FILE_PATH.name}. Please close the Excel file and try again."
    finally:
        workbook.close()


def normalize_values(values):
    return [value.strip() for value in values]


def get_row_values(sheet, row_number):
    return [
        "" if sheet.cell(row=row_number, column=column).value is None else str(sheet.cell(row=row_number, column=column).value)
        for column in range(1, len(HEADERS) + 1)
    ]


def build_updated_values(current_values, values_to_append):
    return [
        current + extra if extra else current
        for current, extra in zip(current_values, normalize_values(values_to_append))
    ]


def validate_contact_values(values):
    name, email, mobile_no, address, branch = normalize_values(values)

    if not name:
        return False, "Name is required."
    if not NAME_PATTERN.fullmatch(name):
        return False, "Name should contain only letters and spaces."

    if not email:
        return False, "Email is required."
    if not EMAIL_PATTERN.fullmatch(email):
        return False, "Enter a valid email address."

    if not mobile_no:
        return False, "Mobile number is required."
    if not MOBILE_PATTERN.fullmatch(mobile_no):
        return False, "Mobile number must contain exactly 10 digits."

    if not address:
        return False, "Address is required."
    if not ADDRESS_PATTERN.fullmatch(address):
        return False, "Address should be 5 to 100 characters and may use letters, numbers, and , . / # -"

    if not branch:
        return False, "Branch is required."
    if not BRANCH_PATTERN.fullmatch(branch):
        return False, "Branch should contain only letters, numbers, spaces, or hyphens."

    return True, ""


def get_all_contacts():
    workbook, sheet = get_workbook_and_sheet()
    contacts = [(row_number, *get_row_values(sheet, row_number)) for row_number in range(2, sheet.max_row + 1)]
    workbook.close()
    return contacts


def create_contact(values):
    cleaned_values = normalize_values(values)
    is_valid, message = validate_contact_values(cleaned_values)
    if not is_valid:
        return False, message

    workbook, sheet = get_workbook_and_sheet()
    sheet.append(cleaned_values)
    return save_workbook(workbook)


def update_contact(row_number, values_to_append):
    cleaned_append_values = normalize_values(values_to_append)
    if not any(cleaned_append_values):
        return False, "Enter at least one value to append."

    workbook, sheet = get_workbook_and_sheet()
    if row_number < 2 or row_number > sheet.max_row:
        workbook.close()
        return False, "Row number not found."

    updated_values = build_updated_values(get_row_values(sheet, row_number), cleaned_append_values)
    is_valid, message = validate_contact_values(updated_values)
    if not is_valid:
        workbook.close()
        return False, message

    for column, value in enumerate(updated_values, start=1):
        sheet.cell(row=row_number, column=column).value = value

    return save_workbook(workbook)


def delete_contact(row_number):
    workbook, sheet = get_workbook_and_sheet()
    if row_number < 2 or row_number > sheet.max_row:
        workbook.close()
        return False, "Row number not found."

    sheet.delete_rows(row_number, 1)
    return save_workbook(workbook)


class ContactFormApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Contact Form CRUD")
        self.root.geometry("920x540")

        self.status_var = tk.StringVar(value="Ready")
        self.row_var = tk.StringVar()
        self.entry_vars = {header: tk.StringVar() for header in HEADERS}

        self.build_ui()
        self.refresh_contacts()

    def build_ui(self):
        main_frame = ttk.Frame(self.root, padding=16)
        main_frame.pack(fill="both", expand=True)

        ttk.Label(main_frame, text="Contact Form", font=("Arial", 16, "bold")).pack(anchor="w")
        ttk.Label(
            main_frame,
            text="Add inserts a new row. Update appends your new text and validates the final result.",
        ).pack(anchor="w", pady=(4, 12))

        form_frame = ttk.Frame(main_frame)
        form_frame.pack(fill="x")

        ttk.Label(form_frame, text="Row Number").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        ttk.Entry(form_frame, textvariable=self.row_var, width=20).grid(
            row=0, column=1, padx=5, pady=5, sticky="w"
        )

        for index, header in enumerate(HEADERS, start=1):
            ttk.Label(form_frame, text=header).grid(row=index, column=0, padx=5, pady=5, sticky="w")
            ttk.Entry(form_frame, textvariable=self.entry_vars[header], width=40).grid(
                row=index, column=1, padx=5, pady=5, sticky="ew"
            )

        form_frame.columnconfigure(1, weight=1)

        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill="x", pady=10)

        ttk.Button(button_frame, text="Add", command=self.handle_add).pack(side="left", padx=5)
        ttk.Button(button_frame, text="Read / Refresh", command=self.refresh_contacts).pack(
            side="left", padx=5
        )
        ttk.Button(button_frame, text="Update", command=self.handle_update).pack(side="left", padx=5)
        ttk.Button(button_frame, text="Delete", command=self.handle_delete).pack(side="left", padx=5)
        ttk.Button(button_frame, text="Clear", command=self.clear_inputs).pack(side="left", padx=5)

        table_frame = ttk.Frame(main_frame)
        table_frame.pack(fill="both", expand=True, pady=(8, 0))

        columns = ["Row Number", *HEADERS]
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings")
        for column in columns:
            self.tree.heading(column, text=column)
            self.tree.column(column, width=130, anchor="w")

        self.tree.bind("<<TreeviewSelect>>", self.on_row_select)
        self.tree.pack(side="left", fill="both", expand=True)

        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        scrollbar.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=scrollbar.set)

        ttk.Label(main_frame, textvariable=self.status_var).pack(anchor="w", pady=(10, 0))

    def get_form_values(self):
        return [self.entry_vars[header].get() for header in HEADERS]

    def clear_inputs(self):
        self.row_var.set("")
        for variable in self.entry_vars.values():
            variable.set("")
        self.status_var.set("Inputs cleared.")

    def refresh_contacts(self):
        self.tree.delete(*self.tree.get_children())
        for contact in get_all_contacts():
            self.tree.insert("", "end", values=contact)
        self.status_var.set("Contact list refreshed.")

    def parse_row_number(self):
        try:
            return int(self.row_var.get())
        except ValueError:
            return None

    def on_row_select(self, _event):
        selected_items = self.tree.selection()
        if not selected_items:
            return

        values = self.tree.item(selected_items[0], "values")
        self.row_var.set(values[0])
        for variable in self.entry_vars.values():
            variable.set("")
        self.status_var.set("Row selected. Type text in the fields to append, then click Update.")

    def handle_add(self):
        success, message = create_contact(self.get_form_values())
        if success:
            self.refresh_contacts()
            self.clear_inputs()
            self.status_var.set(message)
            messagebox.showinfo("Add", "Contact added successfully.")
        else:
            self.status_var.set(message)
            messagebox.showwarning("Add", message)

    def handle_update(self):
        row_number = self.parse_row_number()
        if row_number is None:
            messagebox.showwarning("Update", "Enter a valid row number.")
            return

        success, message = update_contact(row_number, self.get_form_values())
        if success:
            self.refresh_contacts()
            self.clear_inputs()
            self.status_var.set("Record updated.")
            messagebox.showinfo("Update", "Contact updated successfully.")
        else:
            self.status_var.set(message)
            messagebox.showwarning("Update", message)

    def handle_delete(self):
        row_number = self.parse_row_number()
        if row_number is None:
            messagebox.showwarning("Delete", "Enter a valid row number.")
            return

        success, message = delete_contact(row_number)
        if success:
            self.refresh_contacts()
            self.clear_inputs()
            self.status_var.set("Record deleted.")
            messagebox.showinfo("Delete", "Contact deleted successfully.")
        else:
            self.status_var.set(message)
            messagebox.showwarning("Delete", message)


def main():
    root = tk.Tk()
    ContactFormApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()